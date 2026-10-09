import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Set standard margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Title & Organization Header
p_org = doc.add_paragraph()
p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_org = p_org.add_run("मध्यप्रदेश विज्ञान एवं प्रौद्योगिकी परिषद, भोपाल\n")
run_org.font.name = "Arial"
run_org.font.size = Pt(18)
run_org.font.bold = True
run_org.font.color.rgb = RGBColor(15, 52, 96)

run_sub = p_org.add_run("विज्ञान भवन, नेहरू नगर, भोपाल - 462 003 | फोन: 0755-2433138 | www.mpcost.nic.in\n")
run_sub.font.size = Pt(10)
run_sub.font.color.rgb = RGBColor(100, 116, 139)

run_scheme = p_org.add_run("योजना: विज्ञान के प्रचार-प्रसार / विज्ञान लोकव्यापीकरण योजना\n")
run_scheme.font.size = Pt(11)
run_scheme.font.bold = True
run_scheme.font.color.rgb = RGBColor(3, 105, 161)

# Project Title Box Table
tbl_title = doc.add_table(rows=1, cols=1)
tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_title = tbl_title.cell(0, 0)
set_cell_shading(cell_title, "1E293B")
set_cell_margins(cell_title, 150, 150, 200, 200)

p_t = cell_title.paragraphs[0]
p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t1 = p_t.add_run("प्रस्तावित कार्यक्रम: शाला के विद्यार्थियों की विज्ञान प्रतियोगिता\n")
r_t1.font.size = Pt(15)
r_t1.font.bold = True
r_t1.font.color.rgb = RGBColor(56, 189, 248)

r_t2 = p_t.add_run("(3-दिवसीय ग्रामीण विज्ञान संवर्धन कार्यशाला, मॉडल मेकिंग एवं विज्ञान प्रश्नमंच)\n\n")
r_t2.font.size = Pt(11)
r_t2.font.color.rgb = RGBColor(241, 245, 249)

r_t3 = p_t.add_run("लक्षित वर्ग: कक्षा 6वीं से 12वीं के स्कूली विद्यार्थी | अवधि: 3 दिवस | भाषा: हिन्दी / अंग्रेजी\nकुल प्रस्तावित बजट: ₹ 2,39,100/- (दो लाख उनतालीस हजार एक सौ रुपये मात्र)")
r_t3.font.size = Pt(10.5)
r_t3.font.bold = True
r_t3.font.color.rgb = RGBColor(254, 240, 138)

doc.add_paragraph()

# Function for clean headings
def add_custom_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 52, 96)
    return p

# 1. Executive Summary
add_custom_heading("1. प्रस्तावना एवं आवश्यकता (Executive Summary & Local Need)")
p1 = doc.add_paragraph(
    "मध्य प्रदेश के ग्रामीण और दूरदराज के अंचलों में स्कूली विद्यार्थियों के पास प्रतिभा और जिज्ञासा की कोई कमी नहीं है, "
    "परंतु आधुनिक विज्ञान प्रयोगशालाओं (Science Laboratories), हैंड्स-ऑन किट्स और प्रायोगिक उपकरणों के अभाव में उनकी "
    "शिक्षा केवल पाठ्यपुस्तकों और सैद्धांतिक रटने तक सीमित रह जाती है।\n\n"
    "इस अंतर को समाप्त करने के लिए मध्यप्रदेश विज्ञान एवं प्रौद्योगिकी परिषद (MPCST) की 'विज्ञान लोकव्यापीकरण योजना' "
    "के अंतर्गत इस 3-दिवसीय विज्ञान प्रतियोगिता एवं व्यावहारिक कार्यशाला की रूपरेखा तैयार की गई है। इसका उद्देश्य विद्यार्थियों "
    "को 'करके सीखो' (Learning by Doing) की पद्धति से जोड़ना, तार्किक दृष्टिकोण विकसित करना तथा स्थानीय समस्याओं (जैसे सौर ऊर्जा, "
    "जल संचयन और कचरा प्रबंधन) के समाधान हेतु वैज्ञानिक नवाचार को प्रोत्साहित करना है।"
)
p1.paragraph_format.line_spacing = 1.15

# 2. Key Objectives
add_custom_heading("2. परियोजना के मुख्य उद्देश्य (Key Objectives)")
objectives = [
    ("वैज्ञानिक चेतना का विकास", "ग्रामीण बच्चों में अंधविश्वास और झिझक समाप्त कर तार्किक व खोजी सोच पैदा करना।"),
    ("व्यावहारिक मॉडल निर्माण", "प्रकाश, लेंस, दर्पण, ध्वनि, साधारण परिपथ और चुंबकत्व के सिद्धांतों को स्वयं अपने हाथों से मॉडल बनाकर समझना।"),
    ("टीमवर्क एवं नेतृत्व क्षमता", "सामूहिक गतिविधियों और ग्रुप प्रोजेक्ट्स के माध्यम से सहयोग, मंच प्रस्तुति और संवाद कौशल को निखारना।"),
    ("डिजिटल साक्षरता की प्रेरणा", "प्रोजेक्टर, कंप्यूटर और विज्ञान वीडियो के माध्यम से आधुनिक वैज्ञानिक दुनिया और नवाचारों से परिचय कराना।"),
    ("राज्य के विकास में योगदान", "मध्य प्रदेश के सतत विकास लक्ष्यों (SDGs) के अनुरूप भविष्य के नवाचारियों और वैज्ञानिकों की पौध तैयार करना।")
]
for title, desc in objectives:
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(f"{title}: ")
    r1.font.bold = True
    r2 = p.add_run(desc)

# 3. 3-Day Schedule
add_custom_heading("3. तीन-दिवसीय विस्तृत कार्ययोजना (3-Day Execution Schedule)")
tbl_sched = doc.add_table(rows=1, cols=3)
tbl_sched.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = tbl_sched.rows[0].cells
hdr_cells[0].text = "दिवस"
hdr_cells[1].text = "समय"
hdr_cells[2].text = "गतिविधि एवं विवरण"
for c in hdr_cells:
    set_cell_shading(c, "0F3460")
    for r in c.paragraphs[0].runs:
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

sched_data = [
    ("दिवस 1", "09:30 AM - 11:30 AM", "पंजीयन, आईडी कार्ड एवं किट वितरण | दीप प्रज्वलन एवं उद्घाटन सत्र"),
    ("दिवस 1", "11:30 AM - 01:00 PM", "रोचक विज्ञान प्रदर्शन (Fun with Science) | प्रोजेक्टर पर वैज्ञानिक खोजों की डॉक्यूमेंट्री"),
    ("दिवस 1", "02:00 PM - 04:30 PM", "विज्ञान प्रश्नमंच (Science Quiz) | कक्षावार लिखित व ओरल बज़र राउंड"),
    ("दिवस 2", "09:30 AM - 01:00 PM", "हैंड्स-ऑन मॉडल मेकिंग कार्यशाला | प्रकाश, लेंस, दर्पण, मोटर, विद्युत स्विच व तार संयोजन"),
    ("दिवस 2", "02:00 PM - 04:30 PM", "पर्यावरण एवं ऊर्जा नवाचार | सौर ऊर्जा, जल संचयन, कचरा प्रबंधन व चार्ट प्रेजेंटेशन"),
    ("दिवस 3", "09:30 AM - 01:30 PM", "भव्य विज्ञान प्रदर्शनी | बच्चों द्वारा मॉडल्स की मौखिक व्याख्या व निर्णायक मंडल द्वारा मूल्यांकन"),
    ("दिवस 3", "02:30 PM - 04:30 PM", "समापन एवं सम्मान समारोह | विजेताओं को शील्ड/पुरस्कार, सहभागिता प्रमाण-पत्र वितरण एवं आभार")
]

for d, t, a in sched_data:
    row = tbl_sched.add_row()
    row.cells[0].text = d
    row.cells[1].text = t
    row.cells[2].text = a
    for c in row.cells:
        set_cell_margins(c, 60, 60, 100, 100)
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

doc.add_paragraph()

# 4. Budget Breakdown
add_custom_heading("4. प्रस्तावित मदवार बजट विवरणी (Head-wise Itemized Budget)")
tbl_b = doc.add_table(rows=1, cols=3)
tbl_b.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_b = tbl_b.rows[0].cells
hdr_b[0].text = "क्र."
hdr_b[1].text = "मद (Head of Expenditure)"
hdr_b[2].text = "प्रस्तावित राशि (₹)"
for c in hdr_b:
    set_cell_shading(c, "0F3460")
    for r in c.paragraphs[0].runs:
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

budget_items = [
    ("1", "साइंस मॉडल किट्स व प्रदर्शनी सामग्री (Motors, Lenses, Wires, Kits)", "₹ 60,000/-"),
    ("2", "कार्यक्रम आयोजक एवं तकनीकी प्रशिक्षक / रिसोर्स पर्सन मानदेय", "₹ 55,000/-"),
    ("3", "स्टेशनरी, बैकड्रॉप बैनर्स, सहभागिता प्रमाणपत्र, प्रतियोगिता सामग्री", "₹ 40,100/-"),
    ("4", "विज्ञान क्विज एवं अन्य प्रतियोगिता सामग्री व बज़र सिस्टम", "₹ 15,000/-"),
    ("5", "प्रचार-प्रसार सामग्री, पैम्फलेट्स एवं आमंत्रण पत्र", "₹ 15,000/-"),
    ("6", "3 दिवसीय छात्र-छात्राओं, शिक्षकों व अतिथियों हेतु पौष्टिक अल्पाहार/भोजन", "₹ 35,000/-"),
    ("7", "विजेताओं हेतु पुरस्कार, साइंस किट्स एवं अतिथियों हेतु स्मृति चिन्ह", "₹ 19,000/-"),
]

for no, head, amt in budget_items:
    row = tbl_b.add_row()
    row.cells[0].text = no
    row.cells[1].text = head
    row.cells[2].text = amt
    for c in row.cells:
        set_cell_margins(c, 60, 60, 100, 100)
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

row_tot = tbl_b.add_row()
row_tot.cells[0].text = ""
row_tot.cells[1].text = "कुल प्रस्तावित बजट (TOTAL PROPOSED BUDGET):"
row_tot.cells[2].text = "₹ 2,39,100/-"
set_cell_shading(row_tot.cells[0], "E0F2FE")
set_cell_shading(row_tot.cells[1], "E0F2FE")
set_cell_shading(row_tot.cells[2], "E0F2FE")
for c in row_tot.cells:
    set_cell_margins(c, 80, 80, 100, 100)
    for r in c.paragraphs[0].runs:
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(15, 52, 96)

p_amt_words = doc.add_paragraph()
r_w = p_amt_words.add_run("(अक्षरी: दो लाख उनतालीस हजार एक सौ रुपये मात्र)")
r_w.font.italic = True
r_w.font.size = Pt(9.5)
r_w.font.color.rgb = RGBColor(100, 116, 139)

# 5. Video Production Scope
add_custom_heading("5. विशेष कार्ययोजना: वीडियो एवं मल्टीमीडिया निर्माण योजना (Video Production Scope)")

p_v_intro = doc.add_paragraph(
    "प्रस्ताव प्रपत्र के बिंदु 4.12 (मॉनिटरिंग प्रक्रिया) अनुसार:\n"
    "\"हर दिन की गतिविधियों को रजिस्टर में नोट किया जाएगा। कार्यक्रम की तस्वीरें एवं विडियो बनाई जाएगी।\"\n"
    "परिषद में अंतिम रिपोर्ट व उपयोगिता प्रमाण पत्र (UC) प्रेषित करने तथा सोशल मीडिया/प्रचार-प्रसार हेतु निम्नलिखित 4-स्तरीय वीडियो पैकेज प्रस्तावित है:"
)

v_packages = [
    ("1. ऑफिशियल मास्टर डॉक्यूमेंट्री (5 से 7 मिनट)", "संपूर्ण 3-दिवसीय आयोजन का हाई-डेफिनिशन सारांश, छात्र व शिक्षक इंटरव्यू, ग्रामीण बाल नवाचार, निर्णायक प्रतिक्रियाएं व परिषद (MPCST) भोपाल हेतु आधिकारिक सबमिशन।"),
    ("2. इवेंट टीज़र / प्रोमो (45 से 60 सेकंड)", "हाई-एनर्जी म्यूजिक, आधुनिक कट्स और मोशन ग्राफिक्स के साथ सोशल मीडिया और व्हाट्सएप शेयरिंग हेतु आकर्षक प्रोमो।"),
    ("3. डेली सोशल मीडिया रील्स (3 x 60 सेकंड)", "प्रतिदिन शाम को त्वरित एडिट के साथ इंस्टाग्राम, यूट्यूब शॉर्ट्स और स्टेटस हेतु दिनभर की प्रमुख गतिविधियों की रील्स।"),
    ("4. 'नन्हा कलाम' इनोवेशन सीरीज (5 रील्स)", "शीर्ष 5 ग्रामीण बाल वैज्ञानिकों के व्यक्तिगत मॉडल्स पर केंद्रित 45-60 सेकंड के प्रेरणादायक माइक्रो-इंटरव्यूज।")
]

for vt, vd in v_packages:
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(f"{vt}: ")
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(126, 34, 206)
    r2 = p.add_run(vd)

p_v_specs = doc.add_paragraph()
p_v_specs.paragraph_format.space_before = Pt(6)
r_sp1 = p_v_specs.add_run("तकनीकी मानक: ")
r_sp1.font.bold = True
r_sp2 = p_v_specs.add_run("4K/1080p सिनेमैटिक कवरेज, वायरलेस लैपल माइक द्वारा क्रिस्टल क्लियर ऑडियो (जीरो बैकग्राउंड नॉइज), प्रेरणादायक हिन्दी वॉयस-ओवर, एवं MPCST ब्रांडिंग।")

# 6. Checklist
add_custom_heading("6. संलग्न आवश्यक दस्तावेजों की चेकलिस्ट (Compliance Checklist)")
checklist = [
    "परियोजना प्रस्ताव 09 प्रतियों में",
    "विस्तृत मदवार बजट विवरणी",
    "संस्था पंजीयन प्रमाण-पत्र एवं उपनियम/बाय-लॉज प्रति",
    "विगत 3 वर्षों का वार्षिक प्रगति प्रतिवेदन",
    "विगत 3 वर्षों की सीए ऑडिटेड बैलेंस शीट",
    "वर्तमान कार्यकारिणी पदाधिकारियों की धारा 27 अंतर्गत सूची",
    "बैंक खाता विवरण, IFSC कोड एवं निरस्त चेक प्रति",
    "नीति आयोग दर्पण पोर्टल (NGO Darpan) पंजीयन प्रति",
    "अध्यक्ष एवं सचिव के आधार कार्ड की छायाप्रति",
    "गैर-काली सूची (Non-Blacklisting) में होने का शपथ-पत्र",
    "परिषद के नियमों एवं शर्तों की स्वीकृति पत्र",
    "सूचना का अधिकार अधिनियम (RTI 2005) घोषणा पत्र"
]
for item in checklist:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"[ √ ]  {item}")

doc.save(r"D:\GARUDA-AI\docs\MPCST_Project_Report.docx")
print("SUCCESS: D:\\GARUDA-AI\\docs\\MPCST_Project_Report.docx created successfully!")
